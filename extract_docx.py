import sys
from docx import Document

src = r"C:\Users\10679\Desktop\校园任务岛.docx"
out = r"E:\agent1\_docx_dump.txt"

doc = Document(src)
lines = []

def emit(text=""):
    lines.append(text)

# iterate body in order
from docx.oxml.ns import qn
body = doc.element.body
para_map = {p._element: p for p in doc.paragraphs}
table_map = {t._tbl: t for t in doc.tables}

for child in body.iterchildren():
    if child in para_map:
        p = para_map[child]
        style = p.style.name if p.style is not None else ""
        txt = p.text.strip()
        if txt:
            emit(f"[{style}] {txt}")
    elif child in table_map:
        t = table_map[child]
        emit("---- TABLE ----")
        for row in t.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            emit(" | ".join(cells))
        emit("---- END TABLE ----")

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} lines={len(lines)}")
print(f"wrote {out}")
